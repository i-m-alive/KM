"""Detect-and-block check: does a masked surface still appear in a hyperlink
TARGET (the href), not just the visible display text the text-masking pass
already handles? Deliberately detect-only - rewriting hyperlink targets is a
follow-up, not implemented here.

Full package rels traversal is used here deliberately (unlike image
enumeration, where a raw media glob is both simpler and more complete) -
hyperlinks live in relationship files scattered across the package (the main
document, every header, every footer, comments), and package.iter_rels()
sweeps all of them in one pass.
"""

import re

from app.masking.pattern import surface_pattern


def _find_in_text(text: str, surfaces: list[str], where: str) -> list[str]:
    hits = []
    for surface in surfaces:
        if re.search(surface_pattern(surface), text, flags=re.IGNORECASE):
            hits.append(f"{where}: '{surface}'")
    return hits


def _scan_docx(path: str, surfaces: list[str]) -> list[str]:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    hits: list[str] = []
    document = docx.Document(path)
    for rel in document.part.package.iter_rels():
        if rel.reltype == RT.HYPERLINK and rel.is_external and rel.target_ref:
            hits.extend(_find_in_text(rel.target_ref, surfaces, "hyperlink target"))
    return hits


def _scan_pptx(path: str, surfaces: list[str]) -> list[str]:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    hits: list[str] = []
    prs = Presentation(path)
    for rel in prs.part.package.iter_rels():
        if rel.reltype == RT.HYPERLINK and rel.is_external and rel.target_ref:
            hits.extend(_find_in_text(rel.target_ref, surfaces, "hyperlink target"))
    return hits


def _scan_xlsx(path: str, surfaces: list[str]) -> list[str]:
    import openpyxl

    hits: list[str] = []
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink and cell.hyperlink.target:
                    hits.extend(_find_in_text(cell.hyperlink.target, surfaces, f"hyperlink on {ws.title}!{cell.coordinate}"))
    return hits


def _scan_pdf(path: str, surfaces: list[str]) -> list[str]:
    import fitz

    hits: list[str] = []
    doc = fitz.open(path)
    for page_number, page in enumerate(doc):
        for link in page.get_links():
            if link.get("kind") == fitz.LINK_URI and link.get("uri"):
                hits.extend(_find_in_text(link["uri"], surfaces, f"hyperlink on page {page_number + 1}"))
    doc.close()
    return hits


def find_residual_hyperlinks(path: str, content_type: str, filename: str, surfaces: list[str]) -> list[str]:
    if not surfaces:
        return []
    lower = filename.lower()
    if content_type == "application/pdf" or lower.endswith(".pdf"):
        return _scan_pdf(path, surfaces)
    if lower.endswith(".docx") or "wordprocessingml" in content_type:
        return _scan_docx(path, surfaces)
    if lower.endswith(".pptx") or "presentationml" in content_type:
        return _scan_pptx(path, surfaces)
    if lower.endswith(".xlsx") or "spreadsheetml" in content_type:
        return _scan_xlsx(path, surfaces)
    return []
